"""Render the AlpaTrade user guide as Markdown, PowerPoint and PDF.

Layout follows the MMG guides: one page per step, text on the left, screenshot(s) on the right.
Run:  python docs/user_guide/render_user_guide.py   (needs python-pptx, Pillow and LibreOffice)
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE / "screenshots"
NAME = "alpatrade_user_guide"
VERSION = "v0.40.1"
DATE = "12 October 2026"
COMPANY = "Predictive Labs Ltd"
FOOTER = f"AlpaTrade user guide · {VERSION} · {DATE} · © {COMPANY} 2026"

GREEN = RGBColor(0x1F, 0x5D, 0x43)
INK = RGBColor(0x1B, 0x26, 0x21)
MUTED = RGBColor(0x5F, 0x6B, 0x64)
CREAM = RGBColor(0xF6, 0xF4, 0xEE)
PINK = RGBColor(0xE8, 0x45, 0x7E)
LINE = RGBColor(0xDD, 0xD8, 0xCC)
FONT = "Arial"

P1, P2 = "Part 1 · Getting started", "Part 2 · Advanced features"

# Each page: part label, title, subtitle, bullets, images [(file, crop box or None)], layout.
PAGES = [
    dict(kind="cover", title="AlpaTrade user guide",
         subtitle="The chat-first AI trading desk on Alpaca, from sign-in to your first live strategy",
         bullets=["Part 1 · Getting started: sign in, connect Alpaca from the chat, clone the Mag-7 "
                  "buy-the-dip strategy, backtest it and activate it",
                  "Part 2 · Advanced features: every other menu, from the Leaderboard and Hedge Funds "
                  "to the Dashboard, Settings, News and Profile",
                  f"Version {VERSION[1:]} · {DATE} · {COMPANY}"],
         images=[("01-home.png", None)]),
    dict(part=P1, title="1. Sign in and open the chat",
         bullets=["Go to alpatrade.chat and choose Sign in, or Start to create a free account.",
                  "Sign in with Google, or with your email and password.",
                  "You land in the chat. The left menu holds Trade, Chats, Agents, Research and Account. "
                  "Market news is on the right.",
                  "Type a request in plain English, or tap one of the Try a prompt ideas below the box."],
         images=[("02-chat-home.png", None)]),
    dict(part=P1, title="2. Connect Alpaca from the chat",
         bullets=["In the chat box, type connect alpaca. Type connect alpaca live to set up your live "
                  "keys as well.",
                  "Read the heads-up first. AlpaTrade logs in to Alpaca for you and creates new API keys, "
                  "so your old keys stop working at once.",
                  "Enter the email you use for Alpaca, then your Alpaca password. The box hides what you "
                  "type. The password is used once and never saved.",
                  "Type cancel at any time to stop."],
         images=[("03-connect-alpaca-chat.png", None)]),
    dict(part=P1, title="You only type the 2FA code",
         bullets=["AlpaTrade opens Alpaca's login page in a secure browser on its server and fills in your "
                  "email and password for you (masked in these pictures).",
                  "Alpaca then asks for the 6-digit code from your authenticator app.",
                  "Type that code into the AlpaTrade chat straight away, because codes expire in about 30 "
                  "seconds. It is hidden and never stored.",
                  "You don't open Alpaca yourself. The 2FA code is the only thing you type."],
         images=[("04a-alpaca-login-masked.png", (330, 70, 950, 450)),
                 ("04b-alpaca-2fa-prompt.png", (330, 70, 950, 475))], layout="stack"),
    dict(part=P1, title="New keys created, saved and checked",
         bullets=["AlpaTrade creates a fresh paper key pair on your Alpaca dashboard. The key and secret "
                  "are masked here.",
                  "The keys are checked read-only with Alpaca and then stored encrypted. The secret is "
                  "never sent back to your browser.",
                  "The chat confirms: Done! Alpaca is connected.",
                  "Settings > Integrations now shows Alpaca paper as Configured. Press Check at any time: "
                  "a green OK shows your account and its status.",
                  "Prefer to do it by hand? Paste your keys into the same Settings form."],
         images=[("05a-alpaca-paper-keys-masked.png", None), ("05b-settings-check.png", (0, 0, 1440, 560))],
         layout="stack"),
    dict(part=P1, title="3. Pick Mag-7 Buy-the-Dip and clone it",
         bullets=["Open Leaderboard from the left menu.",
                  "Mag-7 Buy-the-Dip · 3-day hold (live) is the live strategy run by Predictive Labs Ltd. It "
                  "buys a Magnificent-7 stock after a dip of 3% or more from its 20-day high. It sells at +8%, "
                  "at −1.5% or after 3 days.",
                  "Each row shows the annualised return, how long the strategy has run, and its alpha vs SPY.",
                  "Press Clone strategy. You get your own private copy, and a backtest starts in a new chat. "
                  "If you clone the same strategy again, AlpaTrade reuses your copy."],
         images=[("06-leaderboard-clone.png", None)]),
    dict(part=P1, title="4. Backtest your copy in the chat",
         bullets=["The backtest uses daily data from January 2016 to today, with fees and slippage included.",
                  "The table compares your strategy with SPY: total return, annualised return, Sharpe ratio, "
                  "maximum drawdown, number of trades and win rate.",
                  "In this run the strategy made +20.0% in total (+1.8% a year), against +360.6% for SPY. "
                  "Sharpe was 0.19, the worst drawdown −36.3%, with 4,952 trades and 49% of them winners.",
                  "Copy CSV and Download CSV export the figures.",
                  "A backtest is hypothetical. It is not investment advice."],
         images=[("07-backtest-metrics.png", None)]),
    dict(part=P1, title="The equity curve vs SPY",
         bullets=["Below the table, a chart shows your strategy and SPY. Both lines start at 100.",
                  "Download PNG saves the chart.",
                  "To try a change, ask in plain English, for example: change the stop to 2% and rerun "
                  "strategy #37. You can also tap a suggested follow-up.",
                  "Open the strategy page to read the full rules or to rerun the backtest."],
         images=[("08-backtest-equity-curve.png", None)]),
    dict(part=P1, title="5. Activate on paper or live",
         bullets=["Under the backtest, the chat shows your status and two buttons.",
                  "Activate Paper trades with paper money (up to $10,000) on your Alpaca paper account until "
                  "you stop it. It starts within a minute and trades during US market hours (ET).",
                  "Activate Live… always opens a confirmation page first. One click never starts live trading.",
                  "The same buttons appear on the Leaderboard and under My strategies."],
         images=[("09-activate-buttons.png", None), ("09b-activate-paper-started.png", (0, 0, 1440, 330))],
         layout="stack"),
    dict(part=P1, title="Activate Live: you confirm first",
         bullets=["Activate Live trades real money on your own Alpaca live account, using cash only.",
                  "If no live account is connected, you see this page. Follow the link to Settings > "
                  "Integrations, or type connect alpaca live in the chat.",
                  "With a live account connected, the page shows the amount the strategy may use and asks "
                  "you to confirm. The strategy starts only after you confirm.",
                  "Later you can Pause it (no new buys; open trades still exit) or Deactivate it."],
         images=[("10-activate-live.png", (0, 0, 1440, 560))]),
    dict(part=P1, title="My strategies and the Dashboard card",
         bullets=["My strategies lists every strategy you own, private ones included, with Shown as, "
                  "annualised return, running period and alpha vs SPY.",
                  "A strategy that is trading shows Paper running, with a Stop paper button next to it.",
                  "The Dashboard has a My strategies card with the same controls and the return since you "
                  "activated it.",
                  "Make public lists your strategy on the Leaderboard. Edit and Delete sit beside it."],
         images=[("11a-my-strategies.png", (0, 0, 1440, 560)), ("11b-dashboard-card.png", (0, 230, 1440, 900))],
         layout="stack"),
    dict(kind="divider", part=P2, title="Part 2 · Advanced features",
         subtitle="Everything beyond your first strategy",
         bullets=["Public site: Platform, Leaderboard, Hedge Funds and Developers.",
                  "Trade: Dashboard, Backtests, Paper runs, Live runs, Live account, Leaderboard, "
                  "My strategies and Daily advisor.",
                  "Agents: Backtest, Validate, Reconcile, Paper Trade and Full Cycle.",
                  "Research: Market Map, Charts, Premarket, Model Analytics, News, Hedge Funds and more.",
                  "Account: Agent Pipeline, Data health, Logging, Settings and Help. Profile is at the bottom "
                  "of the menu."],
         images=[("12-menus.png", None)]),
    dict(part=P2, title="Platform and Developers",
         bullets=["Platform explains the engine behind AlpaTrade: broker links, market data, a backtester and "
                  "a paper-trading loop.",
                  "US equities on Alpaca are live today. Crypto, FX and prediction markets are on the roadmap.",
                  "Developers documents the Agent API at api.alpatrade.chat, with an API reference and a "
                  "Swagger page to try calls.",
                  "Command line: alpatrade backtest paper btd-7dp-05sl-1tp-1d-3m runs the same backtester "
                  "from your terminal. The name spells out the settings, and each run saves a dated "
                  "results folder."],
         images=[("13a-platform.png", (0, 0, 1440, 720)), ("13b-developers-cli.png", (0, 60, 1440, 720))],
         layout="stack"),
    dict(part=P2, title="Leaderboard filters and Copy for ChatGPT / Claude",
         bullets=["Filter by All, Live, Backtest, AlpaTrade, AlpaTrade walk-forward or chatwithtraders.com, "
                  "or search for a trader or strategy.",
                  "Backtests carry a BACKTEST label. Only Live entries show a real track record.",
                  "Open in Grok opens the strategy in Grok, ready to discuss.",
                  "Copy to clipboard copies the strategy as a ready-made skill. Paste it into ChatGPT or "
                  "Claude.",
                  "View more opens the full strategy page."],
         images=[("14a-leaderboard-filters.png", (0, 180, 1440, 640)), ("14b-leaderboard-copy.png", (0, 360, 1440, 900))],
         layout="stack"),
    dict(part=P2, title="Hedge Funds",
         bullets=["See what well-known funds hold and how that portfolio performed.",
                  "Returns are estimated from each fund's 13F filings (its disclosed long US stock positions). "
                  "They are not the fund's reported results.",
                  "Columns show each year since 2021, year to date, the last twelve months, the last twelve "
                  "months vs SPY, and coverage.",
                  "There is also a screener over every 13F filer, the largest managers and recent activist "
                  "filings. Free for everyone once signed in."],
         images=[("15-hedge-funds.png", None)]),
    dict(part=P2, title="Dashboard: paper or live, MTD or YTD",
         bullets=["Choose an account (all accounts, a paper account or a live account) and a period: "
                  "MTD or YTD.",
                  "Tiles show equity, P&L, return, unrealised P&L, cash, net deposits, buying power and the "
                  "annualised return since the account opened.",
                  "Three charts: Performance vs SPY (both start at 100), Drawdown, and Daily returns vs SPY.",
                  "Further down: Daily trading advisor, P&L contributors, strategy rankings, open positions "
                  "and My strategies."],
         images=[("16a-dashboard-ytd.png", (0, 0, 1440, 560)), ("16b-dashboard-charts.png", (0, 0, 1440, 900))],
         layout="stack"),
    dict(part=P2, title="Settings > Integrations",
         bullets=["Alpaca paper and Alpaca live (real money) each have their own keys and status.",
                  "Check sends one read-only request to Alpaca and shows OK with the account and its status.",
                  "Keys are stored encrypted and never sent back to your browser. Saving keys never places "
                  "a trade.",
                  "A live strategy runs only after you press Activate Live and confirm.",
                  "Email reports: choose which daily reports reach your sign-in address."],
         images=[("17-settings-integrations.png", None)]),
    dict(part=P2, title="Profile and Shown as",
         bullets=["Profile shows your email, display name and linked Alpaca accounts, with keys masked. "
                  "You can add another paper account or remove one.",
                  "Shown as is the public name on the Leaderboard and the strategy page, up to 80 characters. "
                  "Set it under My strategies > Edit.",
                  "Tick Public to list a strategy on the Leaderboard. Leave it unticked to keep the strategy "
                  "private."],
         images=[("18a-profile.png", (490, 30, 950, 870)), ("18b-shown-as.png", (300, 20, 1440, 600))],
         layout="pair", weights=(0.36, 0.64)),
    dict(part=P2, title="News and research",
         bullets=["The News panel sits to the right of the chat and the Dashboard. It shows the latest "
                  "headlines with the ticker and its move, grouped by catalyst.",
                  "Filter by ticker, company, event, direction or date. Close the panel with >.",
                  "The Research menu adds Market Map (S&P 500 by sector, coloured by return), Charts, "
                  "Premarket, Model Analytics, News Timing, Historical Research and Market Intel.",
                  "You can also ask in the chat: What's the latest news on TSLA? or Show me a market map."],
         images=[("19a-news-panel.png", (0, 0, 1440, 900)), ("19b-market-map.png", (300, 0, 1440, 560))],
         layout="stack"),
    dict(part=P2, title="Runs, advisor, agents and help",
         bullets=["Backtests lists grid-searched runs that are ready to deploy. You can start from a preset "
                  "such as Dip · mega-cap tech or Momentum.",
                  "Paper runs and Live runs list your sessions. Live account is a read-only view of a linked "
                  "live account.",
                  "Daily advisor reviews your paper accounts each day, 15 minutes after the NYSE close.",
                  "The Agents menu runs Backtest, Validate, Reconcile, Paper Trade or Full Cycle. Agent "
                  "Pipeline shows the automated pipeline.",
                  "Help & shortcuts holds the full command reference."],
         images=[("20a-backtests.png", (0, 0, 1440, 330)), ("20b-help.png", (0, 0, 1440, 900))],
         layout="stack"),
]


def _logo(path: Path) -> Path:
    """The AlpaTrade mark (static/favicon.svg) drawn as a PNG."""
    s = 256
    im = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    k = s / 32
    d.rounded_rectangle([0, 0, s - 1, s - 1], radius=7 * k, fill=(0x1F, 0x5D, 0x43, 255))
    for r, a in ((10.5, 72), (6.5, 115)):
        d.polygon([(16 * k, (16 - r) * k), ((16 + r) * k, 16 * k), (16 * k, (16 + r) * k), ((16 - r) * k, 16 * k)],
                  outline=(255, 255, 255, a), width=int(1.3 * k))
    r = 3.4
    d.polygon([(16 * k, (16 - r) * k), ((16 + r) * k, 16 * k), (16 * k, (16 + r) * k), ((16 - r) * k, 16 * k)],
              fill=(255, 255, 255, 255))
    im.save(path)
    return path


def _crop(name: str, box, tmp: Path) -> Path:
    src = SHOTS / name
    if not box:
        return src
    im = Image.open(src)
    x0, y0, x1, y1 = box
    im = im.crop((x0, y0, min(x1, im.width), min(y1, im.height)))
    out = tmp / f"crop-{src.stem}-{y0}.png"
    im.save(out)
    return out


def _text(slide, left, top, width, height, paras, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (text, size, color, bold, after) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.name, p.font.size, p.font.color.rgb, p.font.bold = FONT, Pt(size), color, bold
        p.space_after = Pt(after)
    return box


def _picture(slide, path: Path, left, top, max_w, max_h):
    w, h = Image.open(path).size
    scale = min(max_w / w, max_h / h)
    pw, ph = w * scale, h * scale
    x, y = left + (max_w - pw) / 2, top + (max_h - ph) / 2
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x - .03), Inches(y - .03),
                                   Inches(pw + .06), Inches(ph + .06))
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.color.rgb = LINE
    frame.shadow.inherit = False
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(pw), Inches(ph))


def _base(prs, logo: Path, n: int, part: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CREAM
    slide.shapes.add_picture(str(logo), Inches(.6), Inches(.38), Inches(.36), Inches(.36))
    _text(slide, 1.02, .36, 3.5, .4, [("AlpaTrade", 15, INK, True, 0)])
    if part:
        _text(slide, 7.3, .4, 5.43, .35, [(part.upper(), 10, GREEN, True, 0)]).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.6), Inches(7.0), Inches(12.13), Inches(.012))
    bar.fill.solid()
    bar.fill.fore_color.rgb = LINE
    bar.line.fill.background()
    _text(slide, .6, 7.05, 10, .3, [(FOOTER, 8, MUTED, False, 0)])
    _text(slide, 11.7, 7.05, 1.03, .3, [(str(n), 8, MUTED, False, 0)]).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    return slide


def build_pptx(out: Path, tmp: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    logo = _logo(tmp / "logo.png")
    for n, pg in enumerate(PAGES, 1):
        kind = pg.get("kind", "page")
        slide = _base(prs, logo, n, pg.get("part", ""))
        bullets = [("•  " + b, 13.5 if kind == "page" else 14, INK, False, 9) for b in pg["bullets"]]
        if kind in ("cover", "divider"):
            fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(.95), Inches(.16), Inches(5.8))
            fill.fill.solid()
            fill.fill.fore_color.rgb = GREEN
            fill.line.fill.background()
            _text(slide, .6, 1.15, 5.1, 2.0, [(pg["title"], 34, GREEN, True, 8), (pg["subtitle"], 14, MUTED, False, 0)])
            _text(slide, .6, 3.25, 5.0, 3.6, bullets)
        else:
            _text(slide, .6, 1.0, 5.1, 1.1, [(pg["title"], 25, GREEN, True, 0)], anchor=MSO_ANCHOR.BOTTOM)
            _text(slide, .6, 2.25, 5.0, 4.6, bullets)
        imgs = [_crop(f, box, tmp) for f, box in pg["images"]]
        L, T, W, H = 5.95, 1.0, 6.78, 5.8
        if len(imgs) == 1:
            _picture(slide, imgs[0], L, T, W, H)
        elif pg.get("layout") == "pair":
            g, x = .2, L
            for im, wt in zip(imgs, pg.get("weights", (.5, .5))):
                _picture(slide, im, x, T, (W - g) * wt, H)
                x += (W - g) * wt + g
        else:
            g = .22
            for i, im in enumerate(imgs):
                _picture(slide, im, L, T + i * (H + g) / 2, W, (H - g) / 2)
    prs.core_properties.title = "AlpaTrade user guide"
    prs.core_properties.author = COMPANY
    prs.save(out)
    return out


def build_pdf(pptx: Path, out: Path, tmp: Path) -> Path:
    profile = tmp / "lo-profile"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                    "--outdir", str(tmp), str(pptx)], check=True, capture_output=True, timeout=300)
    shutil.move(str(tmp / (pptx.stem + ".pdf")), out)
    return out


def build_md(out: Path) -> Path:
    lines = [f"# AlpaTrade user guide", "",
             f"**{PAGES[0]['subtitle']}**  ", f"Version {VERSION[1:]} · {DATE} · {COMPANY}", "",
             "Each section matches one page of the PDF and PowerPoint versions: text on the left, "
             "screenshot on the right. Passwords, 2FA codes, API keys and secrets are masked in every picture.", ""]
    part = None
    for n, pg in enumerate(PAGES, 1):
        if pg.get("kind") == "cover":
            lines += [f"![AlpaTrade home page](screenshots/{pg['images'][0][0]})", ""]
            continue
        if pg.get("part") != part:
            part = pg.get("part")
            if n > 2:
                lines += ['<div style="page-break-before: always"></div>', ""]
            lines += [f"## {part}", ""]
        if pg.get("kind") == "divider":
            lines += [pg["subtitle"] + ".", ""]
        else:
            lines += [f"### {pg['title']}", ""]
        imgs = "".join(f'<img src="screenshots/{f}" width="100%"><br>' for f, _ in pg["images"])
        cell = "\n".join(f"<li>{b}</li>" for b in pg["bullets"])
        lines += ["<table><tr>", f'<td width="42%" valign="top"><ul>\n{cell}\n</ul></td>',
                  f'<td width="58%" valign="top">{imgs}</td>', "</tr></table>", ""]
    lines += ["---", "", f"© {COMPANY} 2026. Backtests are hypothetical and are not investment advice."]
    out.write_text("\n".join(lines) + "\n")
    return out


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        pptx = build_pptx(HERE / f"{NAME}.pptx", tmp)
        print(pptx)
        print(build_pdf(pptx, HERE / f"{NAME}.pdf", tmp))
    print(build_md(HERE / f"{NAME}.md"))
    print(f"{len(PAGES)} pages")
